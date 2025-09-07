# Crear documento Word de muestra
from docx import Document
from docx.shared import Inches

# Crear documento
doc = Document()

# Título principal
doc.add_heading('SEGUNDA ENTREGA - REDISEÑO INSTITUCIONAL', 0)
doc.add_heading('Municipio de Quibdó - Chocó', 1)
doc.add_paragraph('Análisis Integral para la Modernización de la Estructura Administrativa Municipal')
doc.add_paragraph('Fecha: 7 de Septiembre de 2025')

# Sección 1: Diagnóstico General
doc.add_heading('1. DIAGNÓSTICO GENERAL', 1)
doc.add_paragraph('El Municipio de Quibdó presenta una estructura organizacional que requiere modernización urgente. La Ley 617 de 2000 establece las bases para la reestructuración administrativa municipal.')

# Sección 2: Situación Actual
doc.add_heading('2. SITUACIÓN ACTUAL', 1)
doc.add_paragraph('Actualmente, el municipio cuenta con aproximadamente 480 empleados distribuidos en 9 dependencias principales. Los ingresos propios representan el 25% del presupuesto total, mientras que las transferencias nacionales constituyen el 75% restante.')

# Sección 3: Problemas Identificados
doc.add_heading('3. PROBLEMAS IDENTIFICADOS', 1)
problemas = [
    'Duplicación de funciones entre dependencias',
    'Falta de sistemas integrados de información',
    'Dificultades en la gestión del talento humano',
    'Limitaciones en el control financiero',
    'Necesidad de modernización tecnológica'
]

for problema in problemas:
    doc.add_paragraph(f' {problema}', style='List Bullet')

# Sección 4: Propuestas de Rediseño
doc.add_heading('4. PROPUESTAS DE REDISEÑO', 1)
doc.add_paragraph('Se propone una reestructuración que incluya:')
doc.add_paragraph(' Centralización de procesos administrativos', style='List Bullet')
doc.add_paragraph(' Implementación de sistemas digitales', style='List Bullet')
doc.add_paragraph(' Optimización de la planta de personal', style='List Bullet')
doc.add_paragraph(' Fortalecimiento del control interno', style='List Bullet')

# Crear tabla de ejemplo
doc.add_heading('5. ESTRUCTURA PROPUESTA', 1)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'

# Encabezados
header_cells = table.rows[0].cells
header_cells[0].text = 'Dependencia'
header_cells[1].text = 'Empleados Actuales'
header_cells[2].text = 'Empleados Propuestos'

# Datos de ejemplo
dependencias_data = [
    ('Alcaldía', '15', '12'),
    ('Secretaría General', '25', '20'),
    ('Secretaría de Hacienda', '35', '30'),
    ('Secretaría de Planeación', '20', '18'),
    ('Secretaría de Educación', '45', '40')
]

for dep, actual, propuesto in dependencias_data:
    row_cells = table.add_row().cells
    row_cells[0].text = dep
    row_cells[1].text = actual
    row_cells[2].text = propuesto

# Conclusiones
doc.add_heading('6. CONCLUSIONES', 1)
doc.add_paragraph('La implementación del rediseño institucional permitirá al Municipio de Quibdó contar con una estructura más eficiente, moderna y adaptada a las necesidades actuales de la ciudadanía.')

# Guardar documento
doc.save('data/SEGUNDA_ENTREGA_7_SEPTIEMBRE.docx')
print(' Documento Word de muestra creado exitosamente')
